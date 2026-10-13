from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'HueViVu_Tong_hop_cong_nghe_va_tinh_nang.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.65), Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0,0,0)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(6)
d.styles['Normal'].paragraph_format.line_spacing = 1.08
d.styles['Title'].font.size = Pt(25)
for st in d.styles:
    if st.element.pPr is not None:
        for border in list(st.element.pPr.findall(qn('w:pBdr'))):
            st.element.pPr.remove(border)
d.styles['Heading 1'].font.size = Pt(17)
d.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.space_before = Pt(10)
    d.styles[name].paragraph_format.space_after = Pt(6)
header=s.header.paragraphs[0]
header.text='HueViVu  |  Tổng hợp công nghệ và tính năng'
header.style='Normal'
header.runs[0].font.size=Pt(9)
footer=s.footer.paragraphs[0]
footer.alignment=2
footer.add_run('HueViVu  •  13 tháng 10 năm 2026  •  Trang ').font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def p(text): d.add_paragraph(text)
def h(text): d.add_heading(text,1)
def sub(text): d.add_heading(text,2)
def bullet(text): d.add_paragraph(text,'List Bullet')
def page(title): d.add_page_break(); h(title)
def table(headers, rows, widths):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,txt in enumerate(headers): t.rows[0].cells[i].text=txt
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row): cells[i].text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF')); cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'85'); el.set(qn('w:type'),'dxa'); margins.append(el)
            cp.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            cp.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.04
                for run in para.runs: run.font.size=Pt(10.5); run.bold=(ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_paragraph('Tổng hợp công nghệ và tính năng HueViVu','Title')
p('Bản tổng hợp kỹ thuật từ mã nguồn dự án • Ngày 13 tháng 10 năm 2026')
p('HueViVu là ứng dụng web hỗ trợ du lịch Huế, xây dựng bằng Next.js, React và TypeScript, lưu dữ liệu bằng SQLite. Hệ thống kết hợp thuật toán lập lịch A*, ngữ cảnh du lịch và API Gemini hoặc API tương thích Anthropic/Claude. Chưa thấy tích hợp ChatGPT/OpenAI trực tiếp.')
p('Phạm vi: công nghệ, kiến trúc, cơ sở dữ liệu, AI, API, thuật toán, tính năng người dùng, quản trị và vận hành. “Đã có” nghĩa là có phần triển khai trong repo; tài liệu này không xác nhận mọi luồng đã hoạt động trên máy chủ production.')
h('1 Ngôn ngữ và nền tảng')
table(['Thành phần','Công nghệ','Vai trò'],[
('Ngôn ngữ chính','TypeScript','Giao diện, API, thuật toán, nghiệp vụ'),('Ngôn ngữ bổ trợ','JavaScript và Python','Script, Service Worker; dịch vụ AI mở rộng'),('Dữ liệu và trình bày','SQL, HTML JSX, CSS','Truy vấn và cấu trúc giao diện'),('Framework','Next.js 15.5.27','App Router và API phía server'),('Giao diện','React 19','Component và trạng thái tương tác'),('Runtime','Node.js','Chạy Next.js và script'),('Thiết kế','Tailwind CSS 3.4, PostCSS, CSS riêng','Bố cục responsive và hiệu ứng'),('Icon và font','Lucide React, Geist, Geist Mono','Biểu tượng; font lưu trong dự án')],[1.35,2.55,3.2])
sub('Kiến trúc tổng thể')
p('App dùng modular monolith: giao diện và API nghiệp vụ cùng nằm trong Next.js; logic được chia thành module địa điểm, lịch trình, xác thực, thời tiết và AI. Backend dùng Next.js Route Handlers, chưa chuyển sang NestJS.')
p('Luồng chính: trình duyệt → giao diện React → API Next.js → SQLite, thuật toán nội bộ hoặc dịch vụ bên ngoài. Redis/BullMQ và Python/FastAPI là các ranh giới mở rộng đã được chuẩn bị.')

page('2 Cơ sở dữ liệu và lưu trữ')
p('SQLite kết nối qua better-sqlite3. File dữ liệu chính: data/huevivu.db. Hiện chưa dùng PostgreSQL, MySQL, MongoDB, Firebase hoặc Supabase trong luồng lưu trữ chính; PostgreSQL là hướng mở rộng.')
table(['Bảng','Dữ liệu quản lý'],[
('users','Tài khoản, vai trò, khách, sở thích'),('places','Địa điểm, tọa độ, giá, giờ, nguồn và xác minh'),('trips','Lịch trình, ngân sách, trạng thái và chia sẻ'),('trip_state','Tiến độ, hoạt động hiện tại, vị trí và mức mệt'),('chat_messages','Hội thoại theo chuyến đi'),('journal_entries','Nhật ký du lịch'),('trip_likes và trip_saves','Dữ liệu tương tác với chuyến đi'),('feedback và trip_feedback','Phản hồi người dùng và chuyến đi'),('user_events','Sự kiện hành vi'),('training_examples','Ví dụ dữ liệu phục vụ chuẩn bị AI'),('audit_logs','Nhật ký thao tác quản trị'),('data_tasks','Công việc kiểm tra và bổ sung dữ liệu'),('tours','Danh mục tour và combo')],[2.1,5.0])
sub('Kỹ thuật lưu trữ')
bullet('WAL mode, khóa ngoại, prepared statements và transaction cho thao tác cần tính nhất quán.')
bullet('Tạo schema, bổ sung cột và seed dữ liệu bằng mã nguồn. Lịch trình, sở thích và danh sách được lưu dạng JSON trong cột TEXT.')
bullet('Có script sao lưu database và kiểm thử khôi phục.')
bullet('Trình duyệt dùng localStorage cho token, mã phiên và lịch trình offline; Cache API lưu tài nguyên tĩnh.')
sub('Mô hình dữ liệu địa điểm')
p('Ngoài tên, mô tả và tọa độ, dữ liệu có giờ mở cửa, giá, thời lượng, trong nhà, thời điểm phù hợp, mức đông, mức vận động, khả năng tiếp cận, mức phụ thuộc thời tiết, thông tin ẩm thực và liên hệ. Nguồn, người xác minh, thời điểm xác minh và trạng thái xuất bản được theo dõi riêng.')

page('3 AI và nhà cung cấp API')
table(['Nhà cung cấp','Hiện trạng','Vai trò'],[
('Google Gemini','Có mã gọi API trực tiếp','Hội thoại, xử lý yêu cầu AI, ảnh'),('Anthropic Claude hoặc gateway','Có đường gọi /v1/messages','Hội thoại và tùy chỉnh lịch'),('Orimise','Có URL và token cấu hình','Gateway tương thích Anthropic'),('ChatGPT OpenAI','Chưa thấy tích hợp trực tiếp','Chưa có đường gọi OpenAI'),('AI tự huấn luyện','Chưa có pipeline hoàn chỉnh','Có nơi lưu ví dụ và khung mở rộng')],[1.65,2.65,2.8])
sub('Cơ chế chọn provider')
bullet('Có GEMINI_API_KEY và không đặt gateway URL: gọi Gemini Native API generateContent.')
bullet('Có gateway URL: gọi endpoint /v1/messages của gateway. Nhánh Anthropic không có gateway hướng tới api.anthropic.com.')
bullet('AI văn bản mặc định trong mã: gemini-3.6-flash hoặc claude-sonnet-4-6. Vision mặc định: gemini-2.0-flash; có thể đổi qua GEMINI_VISION_MODEL.')
p('Các tên model trên là cấu hình đang ghi trong repo, chưa phải xác nhận tính khả dụng của model với tài khoản API.')
sub('Kỹ thuật AI đã triển khai')
bullet('Prompt có dữ liệu địa điểm, lịch trình và ngữ cảnh người dùng.')
bullet('Một số câu hỏi được trả lời trực tiếp từ catalog và tiến độ chuyến đi trước khi cần gọi LLM.')
bullet('Trích xuất JSON từ phản hồi AI, đặt timeout và có một số đường fallback.')
bullet('Gemini Vision nhận ảnh JPEG dạng base64, kết hợp catalog gần người dùng để hỗ trợ câu trả lời.')
bullet('Kiểm tra thay đổi lịch và bảo vệ các hoạt động đã hoàn thành.')
sub('Lập lịch bằng thuật toán')
p('API tạo lịch chính /api/trips hiện gọi generateAstarTrip trong app, không bắt buộc gọi Gemini hoặc Claude. Repo có hàm tạo lịch bằng LLM nhưng đó không phải đường tạo lịch chính hiện tại.')
sub('RAG và YOLO')
p('Dịch vụ Python/FastAPI có endpoint /v1/rag/search và /v1/vision/detect, nhưng hiện trả danh sách rỗng. Chưa có vector database, model YOLO hoặc pipeline retrieval hoàn chỉnh. AI Trainer quản lý ví dụ dữ liệu, chưa tự fine-tune model.')

page('4 Thuật toán bản đồ và thời tiết')
table(['Kỹ thuật','Ứng dụng'],[
('Tìm kiếm A* và priority queue','Tạo chuỗi hoạt động, ưu tiên trạng thái tìm kiếm'),('Haversine và thời gian di chuyển','Ước lượng khoảng cách và thời gian từ tọa độ'),('Chấm điểm địa điểm','Sở thích, độ phổ biến, xác minh, điểm đã xem hoặc bỏ qua'),('Khung giờ và giờ mở cửa','Xếp ăn uống, tham quan, nghỉ và hoạt động tối'),('Validator và optimizer','Kiểm tra tính khả thi và điều chỉnh lịch'),('Chèn điểm và dynamic replanning','Thêm điểm; đề xuất khi mưa, nóng, mệt hoặc trễ'),('Trip change guard','Bảo vệ tiến độ và kiểm tra thay đổi lịch')],[2.35,4.75])
p('Có quy tắc giới hạn di tích trong ngày, xen kẽ hoạt động nặng và nhẹ, giảm điểm khi di chuyển quá xa. Đây là logic và thuật toán trong app, chưa phải mô hình machine learning đã huấn luyện.')
sub('Bản đồ và chỉ đường')
bullet('Leaflet và React Leaflet hiển thị bản đồ khám phá, bản đồ lịch trình; OpenStreetMap cung cấp lớp bản đồ nền.')
bullet('Browser Geolocation API lấy vị trí khi người dùng cấp quyền.')
bullet('Google Maps JavaScript API hiển thị bản đồ Google khi được cấu hình; Google Routes API tính đường, khoảng cách, thời gian và chỉ dẫn từng bước.')
bullet('Google routing có chế độ đi bộ và lái xe. Valhalla là provider thay thế; nhánh hiện tại hỗ trợ đi bộ.')
p('Development có thể dùng endpoint Valhalla demo. Production cần cấu hình endpoint hoặc bộ khóa Google phù hợp.')
sub('Thời tiết và ngữ cảnh')
p('Open-Meteo cung cấp thời tiết hiện tại tại Huế và dự báo 5 ngày. App phân biệt dữ liệu mới, cũ, không rõ và không có; dữ liệu quá cũ không được dùng để đề xuất đổi lịch. Các tín hiệu giờ, vị trí, thời tiết, sở thích, tiến độ và mức mệt được kết hợp để chọn hành động tiếp theo.')

page('5 Camera giọng nói 3D và offline')
table(['Nhóm','Công nghệ','Chức năng'],[
('Camera','getUserMedia','Lấy ảnh cho Live Guide'),('Nhận giọng nói','SpeechRecognition hoặc webkitSpeechRecognition','Nhập câu hỏi khi trình duyệt hỗ trợ'),('Đọc câu trả lời','Web Speech API speechSynthesis','Đọc nội dung Live Guide'),('API đọc văn bản','Endpoint Google Translate TTS','Âm thanh tiếng Việt qua /api/tts'),('Đồ họa 3D','Three.js và WebGL','Tham quan mô hình Ngọ Môn'),('Tải mô hình','GLTF GLB, GLTFLoader','Tải cảnh và mô hình'),('Điều khiển','OrbitControls và controller đi bộ riêng','Xem tổng thể và di chuyển'),('Chuẩn bị asset','glTF Transform, meshoptimizer, Sharp','Tối ưu mô hình và ảnh'),('Thuyết minh','File MP3','Nghe nội dung ở điểm tham quan'),('PWA và offline','Manifest, Service Worker, Cache API, localStorage','Tài nguyên tĩnh và lịch trình đã lưu')],[1.25,2.7,3.15])
sub('Phạm vi Live Guide')
p('Luồng có camera, câu hỏi văn bản hoặc giọng nói, nhận diện ảnh, kể chuyện, mẹo thực tế và gợi ý bước tiếp theo. Khi thiếu khóa Gemini, app có thể trả gợi ý từ catalog và nêu rõ chưa nhận diện được nội dung ảnh.')
sub('Phạm vi tham quan 3D')
p('Cảnh hiện tập trung vào Ngọ Môn, có điểm dừng kể chuyện, tiến độ, điều khiển chạm và âm thanh. Chưa phải toàn bộ Đại Nội đã được số hóa. Repo khai báo React Three Fiber và Drei; cảnh Ngọ Môn hiện được viết bằng Three.js trực tiếp.')
sub('Phạm vi offline')
p('Service Worker lưu tài nguyên tĩnh và mở trang offline khi điều hướng thất bại. Bản lịch trình lưu trên thiết bị hỗ trợ xem khi mất mạng. AI, thời tiết mới và tính đường vẫn cần kết nối; API không được Service Worker cache như tài nguyên tĩnh.')

page('6 Tính năng người dùng')
table(['Nhóm','Tính năng có trong mã nguồn'],[
('Tài khoản và hồ sơ','Đăng ký, đăng nhập, hồ sơ và danh tính khách'),('Sở thích','Danh mục yêu thích, nhịp đi chậm vừa dày, ưu tiên trong nhà'),('Trang chủ theo ngữ cảnh','Gợi ý theo giờ, thời tiết, vị trí, chuyến đi và mức mệt'),('Khám phá','Tìm kiếm, lọc, danh sách hoặc bản đồ, sắp xếp theo sở thích'),('Chi tiết địa điểm','Ảnh, mô tả, giá, giờ, tọa độ, mẹo, nguồn và xác minh'),('Tạo chuyến đi','1 đến 7 ngày, ngày bắt đầu, ngân sách, phong cách, người đi cùng, điểm xuất phát'),('Quản lý lịch trình','Xem theo ngày, thêm điểm, tùy chỉnh, kiểm tra, tối ưu, lưu và xóa'),('Tiến độ','Bắt đầu chuyến đi, hoạt động hiện tại, hoàn thành, mức mệt'),('Điều chỉnh lịch','Đề xuất khi mưa, nóng, trễ hoặc mệt; xem trước và áp dụng'),('Chat trợ lý','Hỏi địa điểm, tiến độ; yêu cầu thêm, bỏ hoặc dời hoạt động'),('Chỉ đường','Điểm tiếp theo, đường đi, khoảng cách, thời gian và hướng dẫn'),('Live Guide và 3D','Camera, giọng nói, nhận diện ảnh, kể chuyện và tham quan Ngọ Môn'),('Cộng đồng','Xem chuyến đi chia sẻ, thích và mở lịch trình'),('Nhật ký','Nội dung, địa điểm, cảm xúc và trạng thái riêng tư'),('Tour và combo','Danh sách, chi tiết, giá, địa điểm và nội dung tour'),('Phản hồi','Góp ý và đánh giá'),('Trợ giúp khẩn cấp','Trang liên hệ, nút gọi và thông tin cần cung cấp'),('Offline và giới thiệu','Lịch trình đã lưu, PWA, showcase, pitch và demo')],[1.65,5.45])
p('Giới hạn đặt tour: giao diện có nút Đặt ngay, nhưng chưa thấy luồng đặt chỗ và thanh toán hoàn chỉnh. Không nên mô tả app là đã tích hợp thanh toán trực tuyến.')

page('7 Quản trị và API nội bộ')
sub('Quản trị và chất lượng dữ liệu')
bullet('Quản lý người dùng với vai trò user, collector, reviewer và admin.')
bullet('Thêm, sửa, xóa, nhập, xuất địa điểm; phát hiện bản ghi trùng; quản lý tour và phản hồi.')
bullet('Thu thập giờ, giá, tọa độ, liên hệ và đặc điểm trải nghiệm; lưu nguồn và bằng chứng xác minh.')
bullet('Tách trạng thái xuất bản và xác minh; theo dõi dữ liệu cần kiểm tra lại và công việc bổ sung.')
bullet('Audit log ghi thao tác quản trị; AI Trainer quản lý ví dụ dữ liệu.')
sub('Các nhóm API JSON')
table(['Nhóm','Endpoint tiêu biểu'],[
('Xác thực','/api/auth/register, /login, /me, /demo'),('Sở thích','/api/profile/preferences'),('Địa điểm','/api/places, /api/places/[id]; import, export, check-duplicates'),('Chuyến đi','/api/trips, /api/trips/[id]'),('Xử lý lịch','Nhánh trip: validate, optimize, adapt, customize, add-place, place-suggestion, save-itinerary, replan'),('Tiến độ và đường đi','/api/trips/[id]/state, /api/trips/[id]/route'),('Trợ lý','/api/chat, /api/companion'),('Ảnh âm thanh thời tiết','/api/vision, /api/tts, /api/weather'),('Cộng đồng và nhật ký','/api/feed, /api/feed/[id]/like, /api/journal'),('Tour','/api/tours, /api/tours/[id]'),('Phản hồi và hành vi','/api/feedback, /api/events'),('Quản trị','/api/admin/users, /api/admin/data-quality, /api/admin/data-tasks, /api/audit, /api/training'),('Sức khỏe hệ thống','/api/health')],[1.7,5.4])

page('8 Bảo mật vận hành và hướng mở rộng')
sub('Xác thực và bảo mật đã có')
bullet('bcryptjs băm mật khẩu; jsonwebtoken tạo JWT HS256, thời hạn token 30 ngày.')
bullet('Bearer token xác thực API; client lưu token trong localStorage. Kiểm tra vai trò và quyền sở hữu chuyến đi.')
bullet('Kiểm tra đầu vào ở các API quan trọng; UUID tạo định danh.')
bullet('Security headers, request ID, HSTS trong production và audit log.')
sub('Vận hành và kiểm thử')
p('Có logger JSON, /api/health, script kiểm tra biến môi trường production, backup database và kiểm thử restore. Các công cụ kiểm tra gồm TypeScript, ESLint, Next.js build và script Node.js cho lịch trình, ngữ cảnh, sở thích, thời tiết, chỉ đường, trợ lý, offline và 3D.')
p('Lệnh chính: npm run dev để phát triển; npm run build để build; npm run start để chạy bản build; npm run test:types để kiểm tra TypeScript.')
sub('Phần mở rộng và phần chưa hoàn chỉnh')
table(['Công nghệ hoặc tính năng','Trạng thái'],[
('Redis và BullMQ','Có module queue, retry và backoff; chưa thấy worker nghiệp vụ hoàn chỉnh'),('Python FastAPI Pydantic Uvicorn','Có dịch vụ và hợp đồng API; chưa có inference hoặc retrieval thực tế'),('RAG và vector database','Endpoint RAG trả rỗng; chưa tích hợp vector database'),('YOLO','Chưa có model weights và pipeline nhận diện thực tế'),('PostgreSQL','Hướng mở rộng; hiện dùng SQLite'),('NestJS Kafka microservices','Chưa sử dụng'),('Thanh toán VNPay MoMo Stripe','Chưa thấy tích hợp'),('Đa ngôn ngữ','Giao diện và nội dung chủ yếu tiếng Việt')],[2.35,4.75])
p('Các khả năng trên cần được phân biệt với tính năng đã triển khai; không suy ra mức sẵn sàng production chỉ từ thư viện hoặc cấu hình có trong repo.')

page('9 Nguồn đối chiếu trong dự án')
p('Tài liệu được tổng hợp từ mã nguồn và cấu hình của HueViVu tại thời điểm đọc, ngày 13 tháng 10 năm 2026. Phiên bản công nghệ là phiên bản khai báo trong package.json; chưa kiểm chứng khả dụng của model hoặc dịch vụ bằng cuộc gọi API thực tế trong lần xuất tài liệu này.')
table(['Nguồn','Nội dung đối chiếu'],[
('package.json','Dependencies, phiên bản và lệnh chạy kiểm thử'),('docs/ARCHITECTURE.md','Kiến trúc hiện tại và ranh giới mở rộng'),('docs/PROJECT_MAP.md','Entry point và các luồng người dùng'),('src/lib/db.ts','Schema, bảng, migrations, seed và cấu hình SQLite'),('src/lib/auth.ts và middleware.ts','JWT, vai trò và security headers'),('src/lib/ai.ts','Chọn provider, gọi AI, chat và tùy chỉnh lịch'),('src/app/api/trips/route.ts','Luồng tạo lịch chính bằng generateAstarTrip'),('src/lib/astar.ts và itinerary-validator.ts','Lập lịch, kiểm tra và tối ưu'),('src/lib/navigation-route.ts','Google Routes và Valhalla'),('src/lib/hue-weather.ts','Open-Meteo và độ mới dữ liệu'),('src/app/api/vision/route.ts và live/page.tsx','Gemini Vision, camera và giọng nói'),('src/components/3d và src/lib/3d','Cảnh Ngọ Môn và điều khiển di chuyển'),('public/sw.js và offline-trip-store.js','Service Worker và lưu lịch offline'),('services/ai/main.py và requirements.txt','Khung FastAPI, RAG và vision'),('src/lib/jobs/queue.ts','Redis BullMQ và cấu hình job'),('src/app và src/app/api','Trang giao diện, API, quản trị và tính năng')],[3.0,4.1])
sub('Mô tả kỹ thuật có thể dùng để giới thiệu')
p('HueViVu là ứng dụng web du lịch Huế xây dựng bằng Next.js, React và TypeScript, lưu dữ liệu bằng SQLite. Hệ thống kết hợp thuật toán A*, kiểm tra tính khả thi và ngữ cảnh du lịch với Gemini hoặc API tương thích Claude; tích hợp bản đồ, chỉ đường, thời tiết, camera, giọng nói, tham quan 3D và lưu lịch trình offline.')
OUT.parent.mkdir(parents=True,exist_ok=True)
d.core_properties.title='Tổng hợp công nghệ và tính năng HueViVu'
d.core_properties.subject='Kiến trúc công nghệ API và tính năng từ mã nguồn'
d.save(OUT)
print(OUT)
